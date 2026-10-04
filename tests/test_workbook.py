"""The spreadsheet attached to every email."""
from openpyxl import load_workbook

from legotracker.models import PriceRow, Verdict
from legotracker.workbook import build_workbook


def item(set_number="10294", name="Titanic", price=63999.0, shop="official",
         label="GOOD", mrp=63999.0):
    return {"set_number": set_number, "name": name, "price": price,
            "shop": shop, "url": f"https://x.test/{set_number}",
            "verdict": Verdict(label, 17.0, "17% up its range", True),
            "sources": 1, "history_days": 3, "mrp": mrp, "in_stock": True}


def row(date="2026-10-04", shop="official", set_number="10294",
        price=63999.0, suspect=False):
    return PriceRow(date, shop, set_number, price, 63999.0, True,
                    "https://x.test/p", "feed", suspect)


def test_workbook_has_a_today_and_a_history_sheet(tmp_path):
    path = tmp_path / "prices.xlsx"
    build_workbook(path, [item()], [row()], "2026-10-04")
    assert load_workbook(path).sheetnames == ["Today", "History"]


def test_today_sheet_has_one_row_per_watched_set(tmp_path):
    path = tmp_path / "p.xlsx"
    build_workbook(path, [item(), item("42172", "McLaren P1", 29399.0,
                                       "toycra", "FAIR", 41199.0)],
                   [row()], "2026-10-04")
    sheet = load_workbook(path)["Today"]
    names = [sheet.cell(r, 2).value for r in range(2, sheet.max_row + 1)]
    assert names == ["Titanic", "McLaren P1"]


def test_today_sheet_shows_the_verdict_and_price(tmp_path):
    path = tmp_path / "p.xlsx"
    build_workbook(path, [item()], [row()], "2026-10-04")
    sheet = load_workbook(path)["Today"]
    assert sheet.cell(2, 3).value == "GOOD"
    assert sheet.cell(2, 4).value == 63999.0


def test_saving_off_mrp_as_a_formula_not_a_number(tmp_path):
    # The sheet must recalculate if someone edits a price.
    path = tmp_path / "p.xlsx"
    build_workbook(path, [item(price=29399.0, mrp=41199.0)], [row()],
                   "2026-10-04")
    sheet = load_workbook(path)["Today"]
    header = [c.value for c in sheet[1]]
    cell = sheet.cell(2, header.index("Off MRP") + 1).value
    assert isinstance(cell, str) and cell.startswith("=")


def test_history_sheet_carries_every_recorded_row(tmp_path):
    path = tmp_path / "p.xlsx"
    build_workbook(path, [item()],
                   [row(date="2026-10-03"), row(date="2026-10-04")],
                   "2026-10-04")
    sheet = load_workbook(path)["History"]
    assert sheet.max_row == 3        # header + 2


def test_history_sheet_names_the_set_not_just_its_number(tmp_path):
    path = tmp_path / "p.xlsx"
    build_workbook(path, [item()], [row()], "2026-10-04")
    sheet = load_workbook(path)["History"]
    header = [c.value for c in sheet[1]]
    assert "Name" in header
    assert sheet.cell(2, header.index("Name") + 1).value == "Titanic"


def test_a_suspect_row_is_visibly_flagged(tmp_path):
    path = tmp_path / "p.xlsx"
    build_workbook(path, [item()], [row(price=2999.0, suspect=True)],
                   "2026-10-04")
    sheet = load_workbook(path)["History"]
    header = [c.value for c in sheet[1]]
    assert sheet.cell(2, header.index("Suspect") + 1).value == "YES"


def test_newest_history_first(tmp_path):
    path = tmp_path / "p.xlsx"
    build_workbook(path, [item()],
                   [row(date="2026-10-03"), row(date="2026-10-05")],
                   "2026-10-05")
    sheet = load_workbook(path)["History"]
    assert sheet.cell(2, 1).value == "2026-10-05"


def test_headers_are_frozen_so_scrolling_keeps_them_visible(tmp_path):
    path = tmp_path / "p.xlsx"
    build_workbook(path, [item()], [row()], "2026-10-04")
    book = load_workbook(path)
    assert book["Today"].freeze_panes == "A2"
    assert book["History"].freeze_panes == "A2"


def test_non_ascii_names_survive(tmp_path):
    path = tmp_path / "p.xlsx"
    build_workbook(path, [item(name="RIVENDELL™")], [row()],
                   "2026-10-04")
    assert load_workbook(path)["Today"].cell(2, 2).value == "RIVENDELL™"


def test_it_copes_with_no_history_at_all(tmp_path):
    path = tmp_path / "p.xlsx"
    build_workbook(path, [item()], [], "2026-10-04")
    assert load_workbook(path)["History"].max_row == 1   # header only


def test_history_only_covers_sets_you_watch(tmp_path):
    # The CSV keeps the whole catalogue for the retirement signal, but a
    # sheet with 1,661 rows of sets you do not own is not readable.
    path = tmp_path / "p.xlsx"
    build_workbook(path, [item()],
                   [row(set_number="10294"),
                    row(set_number="99999", price=500.0)],
                   "2026-10-04")
    sheet = load_workbook(path)["History"]
    sets = {sheet.cell(r, 2).value for r in range(2, sheet.max_row + 1)}
    assert sets == {"10294"}


def test_the_off_mrp_formula_points_at_the_right_cells(tmp_path):
    path = tmp_path / "p.xlsx"
    build_workbook(path, [item(price=29399.0, mrp=41199.0)], [row()],
                   "2026-10-04")
    sheet = load_workbook(path)["Today"]
    formula = sheet.cell(2, 6).value
    assert "G2" in formula and "D2" in formula      # MRP and Best price
    assert sheet.cell(2, 4).value == 29399.0        # D2 is the price
    assert sheet.cell(2, 7).value == 41199.0        # G2 is the MRP


def test_running_twice_in_a_day_does_not_double_the_history_rows(tmp_path):
    path = tmp_path / "p.xlsx"
    build_workbook(path, [item()],
                   [row(date="2026-10-04"), row(date="2026-10-04")],
                   "2026-10-04")
    sheet = load_workbook(path)["History"]
    assert sheet.max_row == 2      # header + one row, not two


def test_a_rerun_with_a_changed_price_keeps_the_newer_one(tmp_path):
    path = tmp_path / "p.xlsx"
    build_workbook(path, [item()],
                   [row(date="2026-10-04", price=63999.0),
                    row(date="2026-10-04", price=59999.0)],
                   "2026-10-04")
    sheet = load_workbook(path)["History"]
    assert sheet.max_row == 2
    assert sheet.cell(2, 5).value == 59999.0
